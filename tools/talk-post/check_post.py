#!/usr/bin/env python3
"""Checks for the generated talk post. Run from the repo root.

    python tools/talk-post/check_post.py
        the post is what build_post.py produces, and carries the transcript's text line for line

    python tools/talk-post/check_post.py --html _site/2026/09/18/index.html
        also check the built page (needs beautifulsoup4 and lxml): every highlighted area has its
        text element and the reverse, table cells get their keys, headings keep their anchors

    python tools/talk-post/check_post.py --deck path/to/deck.pptx
        also check the transcript against the deck, paragraph by paragraph (needs python-pptx).
        Reads slide text only, never the speaker notes. Keep the deck outside the repo.
"""
import argparse
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_post as bp  # noqa: E402

failures = []


def report(ok, label, detail=""):
    print(("ok    " if ok else "FAIL  ") + label + (f": {detail}" if detail and not ok else ""))
    if not ok:
        failures.append(label)


def text_units(md):
    """The text lines of a post, with everything build_post.py adds taken away again."""
    body = md[md.index("\n---\n", 4) + 5:]
    body = re.sub(r"<(style|script)>.*?</\1>", "", body, flags=re.S)
    out = []
    for line in body.split("\n"):
        s = re.sub(r"^(>\s?)+", "", line.strip()).strip()
        if not s or s == "---" or bp.IAL.match(s) or s.startswith("<div") or s == "</div>":
            continue
        if s.startswith("!["):
            s = re.match(r"^!\[.*?\]\(.*?\)", s).group(0)
        s = re.sub(r"^(- |\d+\. )\{:[^}]*\}\s*", r"\1", s)
        out.append(re.sub(r"</?span[^>]*>", "", s))
    return out


def check_source():
    transcript, regions, figure_regions, css, js = bp.load()
    post, rows, links = bp.build(transcript, regions, figure_regions, css, js)
    on_disk = bp.POST.read_bytes().decode("utf-8").replace("\r\n", "\n")
    report(on_disk == post, "post matches what build_post.py generates")
    a, b = text_units(transcript), text_units(on_disk)
    first = next((i for i, (x, y) in enumerate(zip(a, b)) if x != y), None)
    report(a == b, f"post carries the transcript's {len(a)} text lines unchanged and in order",
           "" if a == b else (f"first difference at line {first}: {a[first][:60]!r} / {b[first][:60]!r}" if first is not None else f"{len(a)} vs {len(b)} lines"))
    unused = [(r["slide"], t) for r in rows for t in r["unlinked"]]
    report(not unused, "every saved region is used", ", ".join(f"slide {n}: {t[:40]!r}" for n, t in unused))
    return links


QUOTES = {ord("\u201c"): '"', ord("\u201d"): '"', ord("\u2018"): "'", ord("\u2019"): "'"}


def norm(s):
    s = unicodedata.normalize("NFC", s).translate(QUOTES).replace("\u00a0", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return re.sub(r"^\u2014 ", "", s).replace("*", "")


def check_html(path, links):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(Path(path).read_text(encoding="utf-8"), "lxml")
    content = soup.select_one("div.content")
    blocks = content.select("div.sl")
    report(len(blocks) == 31, "31 slide blocks", str(len(blocks)))
    good = 0
    for b in blocks:
        p = b.select_one(".sl-fig > p")
        a = p.find("a", recursive=False) if p else None
        img = a.find("img") if a else None
        good += bool(img is not None and "popup" in a.get("class", []) and a.get("href") == img.get("src") and img.get("alt"))
    report(good == len(blocks), "each slide image has alt text and opens full size", f"{good} of {len(blocks)}")
    order = [re.search(r"-(\d\d)\.png", b.select_one(".sl-fig img")["src"]).group(1) for b in blocks]
    report(order == ["%02d" % n for n in range(2, 33)], "slides 02 to 32 in order")
    problems, areas, elements = [], 0, 0
    for i, b in enumerate(blocks):
        n = str(i + 2)
        for table in b.select("table[data-keys]"):
            keys, cells = table["data-keys"].split(" "), table.select("th, td")
            if len(keys) != len(cells):
                problems.append(f"slide {n}: table has {len(cells)} cells but {len(keys)} keys")
            for cell, key in zip(cells, keys):  # what script.js does in the browser
                if key != "-":
                    cell["data-k"] = key
        regs = b.select(".sl-fig .sl-r")
        by_key = {}
        for e in b.select(".sl-txt [data-k]"):
            by_key.setdefault(e["data-k"], []).append(e)
        head = b.find_previous_sibling()
        if head is not None and head.name == "h2" and head.get("data-k"):
            by_key.setdefault(head["data-k"], []).append(head)
        areas += len(regs)
        elements += sum(len(v) for v in by_key.values())
        region_keys = {r["data-k"] for r in regs}
        problems += [f"slide {n}: area {k} has no text element" for k in region_keys - set(by_key)]
        problems += [f"slide {n}: text element {k} has no area" for k in set(by_key) - region_keys]
        for key, texts in links.get(n, {}).items():
            if key not in by_key:
                problems.append(f"slide {n}: key {key} missing from the page")
                continue
            if len(by_key[key]) != 1:
                problems.append(f"slide {n}: key {key} is on {len(by_key[key])} elements")
            shown = norm(by_key[key][0].get_text())
            for t in texts:
                if norm(t) not in shown and not re.fullmatch(r"\d", norm(t)):
                    problems.append(f"slide {n}: key {key} does not show {norm(t)[:50]!r}")
    report(not problems, f"{elements} text elements and {areas} slide areas are linked consistently", "; ".join(problems[:6]))
    heads = content.select("h2, h3")
    report(all(h.get("id") and h.select_one("a.anchor") for h in heads), f"all {len(heads)} headings keep their id and anchor link")
    scripts = content.select("script")
    report(len(scripts) == 1 and len(content.select("style")) == 1, "one inline style block and one inline script")
    if scripts and shutil.which("node"):
        with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False, encoding="utf-8") as f:
            f.write(scripts[0].get_text())
        r = subprocess.run(["node", "--check", f.name], capture_output=True, text=True)
        Path(f.name).unlink()
        report(r.returncode == 0, "the script still parses after the production build joins its lines", r.stderr.strip()[:160])
    for t in content.select("style, script"):
        t.decompose()
    text = content.get_text()
    left = [tok for tok in ("{:", "data-k", "markdown=", "<span", "{%", "*[") if tok in text]
    report(not left, "no markup has leaked into the visible text", " ".join(left))


def check_deck(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def walk(shapes):
        for s in shapes:
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(s.shapes)
            else:
                yield s

    def paragraphs(slide):
        for s in walk(slide.shapes):
            frames = []
            if getattr(s, "has_table", False) and s.has_table:
                frames = [c.text_frame for r in s.table.rows for c in r.cells]
            elif getattr(s, "has_text_frame", False) and s.has_text_frame:
                frames = [s.text_frame]
            for frame in frames:
                for p in frame.paragraphs:
                    if p.text.strip():
                        yield p.text.replace("\v", " ")

    transcript = bp.load()[0]
    _, sections = bp.split_post(transcript)
    slides = list(Presentation(path).slides)
    report(len(slides) == 32, "the deck has 32 slides", str(len(slides)))
    total, missing = 0, []
    for n, slide in enumerate(slides, 1):
        section = "\n".join(sections[n - 1])
        section = re.sub(r"\{%.*?title='([^']*)'.*?%\}", r"\1", section)  # the video's caption lives in its include
        have = bp.na(re.sub(r"<[^>]+>|^\s*(>\s?)+|[|]", " ", section, flags=re.M))
        have = re.sub(r"\s+", " ", have)
        for p in paragraphs(slide):
            total += 1
            if bp.na(p) not in have:
                missing.append((n, bp.na(p)))
    print(f"      {total} paragraphs of slide text read from the deck")
    for n, t in missing:
        print(f"      slide {n}: not in the transcript: {t[:70]!r}")
    # Expected leftovers: slide 12's tier tag (deliberately not transcribed) and slide 29's numerals 1-4
    # (they are the list numbering). Anything else is a difference between the deck and the transcript.
    unexpected = [m for m in missing if not (m[0] == 12 and len(m[1]) < 24) and not (m[0] == 29 and re.fullmatch(r"\d", m[1]))]
    report(not unexpected, "every other deck paragraph is in the transcript, word for word", f"{len(unexpected)} unexpected")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--html", help="built page to check, e.g. _site/2026/09/18/index.html")
    ap.add_argument("--deck", help="the talk's .pptx (keep it outside the repo)")
    args = ap.parse_args()
    links = check_source()
    if args.html:
        check_html(args.html, links)
    if args.deck:
        check_deck(args.deck)
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
