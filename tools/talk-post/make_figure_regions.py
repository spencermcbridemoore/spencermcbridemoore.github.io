#!/usr/bin/env python3
"""Find where the transcript's "text in the figure" lines sit inside the figure images -> figure-regions.json

    python tools/talk-post/make_figure_regions.py --deck path/to/deck.pptx --work path/to/scratch-folder

Run from the repo root, after make_regions.py. Slides 7, 9, 12 and 16 show a picture that carries
its own text, which the deck's layout knows nothing about. This script takes each picture out of
the deck (into the scratch folder, outside the repo), has Windows recognise the words and their
positions (ocr_figures.ps1), and matches every transcript line that is not native slide text to a
run of recognised words.

This is best effort: recognition misreads characters, so lines are compared loosely, and a line
the transcript phrases in its own words (a description of the figure, an added column heading) is
simply not found and stays unlinked. LOOK AT work/overlays/fig*.png afterwards.
Needs python-pptx, numpy, opencv-python, pillow and Windows (for the OCR step).
"""
import argparse
import difflib
import io
import json
import re
import subprocess
import sys
import unicodedata
from pathlib import Path

import cv2
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_post as bp  # noqa: E402

W, H = 1920, 1080
IMAGE = "assets/img/posts/2026-09-18-ai-environmental-impact-%02d.png"
FIGURE_SLIDES = (7, 9, 12, 16)
LABELS = ("text in the figure", "image captions", "photo caption")  # structure added by the transcript


def key(s):
    """Loose comparison key: recognition confuses l/I and 0/O and drops punctuation."""
    s = unicodedata.normalize("NFKC", s).lower().replace("l", "i").replace("0", "o")
    return re.sub(r"[^a-z0-9]", "", s)


def units_of(lines):
    """The visible text lines of one slide's section of the transcript."""
    out = []
    for line in lines:
        s = line.strip()
        if not s or s.startswith("![") or s.startswith("{:") or s.startswith("{%"):
            continue
        s = re.sub(r"^(>\s?)+", "", s).strip()
        if not s:
            continue
        if s.startswith("|"):
            if bp.TABLE_RULE.fullmatch(s):
                continue
            for cell in s.strip()[1:-1].split("|"):
                out.extend(bp.plain(part) for part in re.split(r"<br\s*/?>", cell) if bp.plain(part))
            continue
        s = re.sub(r"^#{1,6}\s+", "", s)
        s = re.sub(r"^(-|\d+\.)\s+", "", s)
        if bp.plain(s):
            out.append(bp.plain(s))
    return out


def extract_figures(deck, folder):
    """Save each figure slide's largest picture at double size; return its place on the slide."""
    prs = Presentation(str(deck))
    meta = {}
    for n in FIGURE_SLIDES:
        slide = prs.slides[n - 1]
        pic = max((s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE), key=lambda s: s.width * s.height)
        im = Image.open(io.BytesIO(pic.image.blob)).convert("RGB")
        scale = 2 if max(im.size) * 2 <= 9000 else 1  # small print reads better enlarged; OCR accepts up to 10000px
        big = im.resize((im.width * scale, im.height * scale), Image.LANCZOS)
        big.save(folder / f"fig{n:02d}.png")
        meta[n] = {"box": [pic.left / prs.slide_width, pic.top / prs.slide_height, pic.width / prs.slide_width, pic.height / prs.slide_height], "size": big.size}
    return meta


def load_words(path):
    """Recognised words grouped into line fragments (split at wide gaps), each knowing the fragment below it."""
    lines = {}
    for row in path.read_text(encoding="utf-8").splitlines():
        if not row.strip():
            continue
        li, text, x, y, w, h = row.split("\t")
        lines.setdefault(int(li), []).append({"t": text, "x0": float(x), "y0": float(y), "x1": float(x) + float(w), "y1": float(y) + float(h)})
    frags = []
    for li in sorted(lines):
        ws = lines[li]
        cur = [ws[0]]
        hmed = float(np.median([w["y1"] - w["y0"] for w in ws]))
        for w in ws[1:]:
            if w["x0"] - cur[-1]["x1"] > 3.0 * hmed:
                frags.append(cur)
                cur = [w]
            else:
                cur.append(w)
        frags.append(cur)
    F = [{"ws": ws, "x0": min(w["x0"] for w in ws), "x1": max(w["x1"] for w in ws), "y0": min(w["y0"] for w in ws), "y1": max(w["y1"] for w in ws)} for ws in frags]
    for f in F:
        f["h"] = f["y1"] - f["y0"]
    for f in F:  # the fragment directly below, in the same column
        best = None
        for g in F:
            if g is f or g["y0"] <= f["y0"] + 0.4 * f["h"]:
                continue
            gap = g["y0"] - f["y1"]
            if gap > 1.3 * max(f["h"], g["h"]):
                continue
            overlap = min(f["x1"], g["x1"]) - max(f["x0"], g["x0"])
            if overlap < 0.3 * min(f["x1"] - f["x0"], g["x1"] - g["x0"]) and abs(g["x0"] - f["x0"]) > 1.5 * f["h"]:
                continue
            score = (gap, abs(g["x0"] - f["x0"]))
            if best is None or score < best[0]:
                best = (score, g)
        f["next"] = best[1] if best else None
    return F


def chain_from(F, fi, wi, need):
    """Words from word wi of fragment fi onward, continuing down the column until `need` characters are covered."""
    out, f, i, total, hops = [], F[fi], wi, 0, 0
    while f is not None and total < need + 12 and hops < 14:
        for w in f["ws"][i:]:
            out.append((w, f))
            total += len(key(w["t"]))
            if total >= need + 12:
                break
        f, i, hops = f["next"], 0, hops + 1
    return out


def strict_ratio(a, b):
    """Similarity counting only runs of 4+ matching characters, so stray letters cannot pull a match into the next line."""
    m = sum(blk.size for blk in difflib.SequenceMatcher(None, a, b, autojunk=False).get_matching_blocks() if blk.size >= 4)
    return 2.0 * m / max(1, len(a) + len(b))


def centre(ws):
    return np.mean([(w["x0"] + w["x1"]) / 2 for w in ws]), np.mean([(w["y0"] + w["y1"]) / 2 for w in ws])


def best_match(F, text, used, prev_centre, whole_fragment_only=False):
    """(score, words) of the run of recognised words most like `text`, or None."""
    T = key(text)
    if len(T) < 2:
        return None
    cands = []
    for fi, f in enumerate(F):
        for wi in ([0] if whole_fragment_only else range(len(f["ws"]))):
            first = key(f["ws"][wi]["t"])
            if not first:
                continue
            if not whole_fragment_only and difflib.SequenceMatcher(None, T[:max(4, len(first))], first, autojunk=False).ratio() < 0.5:
                continue
            seq = chain_from(F, fi, wi, len(T))
            acc, best = "", None
            for j, (w, wf) in enumerate(seq):
                acc += key(w["t"])
                if len(acc) > len(T) + 8:
                    break
                if len(acc) < 0.6 * len(T):
                    continue
                if whole_fragment_only and w is not wf["ws"][-1]:
                    continue
                r = strict_ratio(T, acc) if len(T) >= 12 else difflib.SequenceMatcher(None, T, acc, autojunk=False).ratio()
                if best is None or r > best[0]:
                    best = (r, j)
            if best is None:
                continue
            r, j = best
            ws = [w for w, _ in seq[:j + 1]]
            if any(id(w) in used for w in ws):
                r -= 0.25
            if wi == 0:
                r += 0.004  # prefer starting at the beginning of a line
            cands.append((r, ws))
    if not cands:
        return None
    top = max(c[0] for c in cands)
    near = [c for c in cands if c[0] >= top - 0.015]

    def dist(c):  # among equally good matches, the one nearest the previous line found
        if prev_centre is None:
            return 0
        cx, cy = centre(c[1])
        return abs(cy - prev_centre[1]) * 2 + abs(cx - prev_centre[0])

    near.sort(key=lambda c: (dist(c), -c[0]))
    return near[0]


def cluster(groups):
    """Merge groups of words that sit on the same or adjacent lines."""
    groups = [list(g) for g in groups]
    merged = True
    while merged:
        merged = False
        for i in range(len(groups)):
            for j in range(i + 1, len(groups)):
                a, b = groups[i], groups[j]
                ah = np.median([w["y1"] - w["y0"] for w in a])
                bh = np.median([w["y1"] - w["y0"] for w in b])
                ay0, ay1 = min(w["y0"] for w in a), max(w["y1"] for w in a)
                by0, by1 = min(w["y0"] for w in b), max(w["y1"] for w in b)
                ax0, ax1 = min(w["x0"] for w in a), max(w["x1"] for w in a)
                bx0, bx1 = min(w["x0"] for w in b), max(w["x1"] for w in b)
                vgap, hgap = max(by0 - ay1, ay0 - by1, 0), max(bx0 - ax1, ax0 - bx1, 0)
                if vgap <= 1.2 * max(ah, bh) and hgap <= 6 * max(ah, bh):
                    groups[i] = a + b
                    del groups[j]
                    merged = True
                    break
            if merged:
                break
    return groups


def match_slide(n, units, F, place):
    bx, by, bw, bh = place["box"]
    iw, ih = place["size"]

    def to_slide(ws):
        x0, x1 = min(w["x0"] for w in ws) / iw, max(w["x1"] for w in ws) / iw
        y0, y1 = min(w["y0"] for w in ws) / ih, max(w["y1"] for w in ws) / ih
        px, py = 7 / W, 5 / H
        return [round(max(0, bx + x0 * bw - px), 5), round(max(0, by + y0 * bh - py), 5), round(min(1, bx + x1 * bw + px), 5), round(min(1, by + y1 * bh + py), 5)]

    used, prev, found, log = set(), None, [], []
    for u in units:
        T = key(u)
        short = len(T) < 9  # a short label must be a whole recognised fragment, or it matches inside sentences
        threshold = 0.93 if short else (0.86 if len(T) < 25 else 0.8)
        hit = best_match(F, u, used, prev, whole_fragment_only=short)
        if hit and hit[0] >= threshold:
            used.update(id(w) for w in hit[1])
            found.append({"text": u, "box": to_slide(hit[1])})
            prev = centre(hit[1])
            continue
        # a line that joins several labels of the figure: look for the pieces separately
        parts = [p.strip() for p in re.split(r" \u00b7 |: |; | \u2014 ", u) if len(key(p)) >= 8]
        groups = []
        if len(parts) >= 2:
            for part in parts:
                h2 = best_match(F, part, used, prev)
                if h2 and h2[0] >= 0.88:
                    used.update(id(w) for w in h2[1])
                    groups.append(h2[1])
                    prev = centre(h2[1])
        if hit and hit[0] >= threshold - 0.12 and not groups:  # a near miss on the whole line is still the right place
            groups.append(hit[1])
            used.update(id(w) for w in hit[1])
        found.extend({"text": u, "box": to_slide(g)} for g in cluster(groups))
        log.append((u, f"{len(groups)} of {len(parts)} pieces" if groups else "NOT FOUND"))
    return found, log


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--deck", required=True, help="the talk's .pptx; keep it outside the repo")
    ap.add_argument("--work", required=True, help="scratch folder outside the repo")
    ap.add_argument("--reuse-ocr", action="store_true", help="use the fig*.tsv files already in the scratch folder")
    args = ap.parse_args()
    work = Path(args.work).resolve()
    if bp.ROOT == work or bp.ROOT in work.parents:
        sys.exit("--work must be outside the repo")
    ocr = work / "ocr"
    ocr.mkdir(parents=True, exist_ok=True)
    (work / "overlays").mkdir(exist_ok=True)
    places = extract_figures(args.deck, ocr)
    if not args.reuse_ocr:
        subprocess.run(["powershell", "-NoProfile", "-File", str(bp.HERE / "ocr_figures.ps1"), "-Dir", str(ocr)], check=True)
    _, sections = bp.split_post((bp.HERE / "transcript.md").read_text(encoding="utf-8"))
    native = json.loads((bp.HERE / "regions.json").read_text(encoding="utf-8"))
    out = {}
    for n in FIGURE_SLIDES:
        native_texts = {bp.na(e["text"]) for e in native.get(str(n), [])}
        units = [u for u in units_of(sections[n - 1]) if u not in native_texts and u.rstrip(":").lower() not in LABELS]
        found, log = match_slide(n, units, load_words(ocr / f"fig{n:02d}.tsv"), places[n])
        out[str(n)] = found
        missing = sum(1 for _, how in log if how == "NOT FOUND")
        print(f"slide {n}: {len(units)} figure-text lines, {len(units) - missing} located, {missing} not")
        for u, how in log:
            print(f"     [{how}] {u[:90]}")
        img = np.array(Image.open(bp.ROOT / (IMAGE % n)).convert("RGB"))
        colours = [(230, 60, 60), (40, 120, 230), (30, 160, 70), (220, 140, 0), (150, 60, 200)]
        for i, e in enumerate(found):
            b = e["box"]
            cv2.rectangle(img, (int(b[0] * W), int(b[1] * H)), (int(b[2] * W), int(b[3] * H)), colours[i % 5], 3)
        Image.fromarray(img).save(work / "overlays" / f"fig{n:02d}.png")
    bp.dump_regions(bp.HERE / "figure-regions.json", out)
    print(f"wrote figure-regions.json: {sum(len(v) for v in out.values())} areas. Now look at {work / 'overlays'} (fig*.png).")


if __name__ == "__main__":
    main()
